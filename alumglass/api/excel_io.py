"""Thư viện dùng chung: tải file mẫu / nạp / xuất Excel cho dialog trong alumglass.

Thiết kế: docs/design/excel-io-shared-library.md

Vì sao gom lại một chỗ: luồng "tải file mẫu -> chọn file -> kiểm tra -> đổ vào bảng" từng bị
chép tay ở nhiều dialog rồi trôi khác nhau — bản Purchase Invoice kiểm header đúng ô A1 còn bản
ĐNTT Domestic Shipment thì không; bản này báo lỗi kèm số dòng, bản kia bỏ qua im lặng; bản này
kiểm quyền trước khi đọc file, bản kia không. Gom về một chỗ thì các luật an toàn chỉ phải sửa
một lần, và dialog mới chỉ còn khai danh sách cột + hàm kiểm tra.

Hai đầu của hợp đồng:
  - DIALOG khai `columns` — MỘT lần, dùng chung cho cả file mẫu, nạp file và xuất Excel — và
    `validator` (dotted path tới hàm Python, tuỳ chọn; không khai thì chạy chế độ raw).
  - VALIDATOR: `def validate_xxx(row, ctx)` trả về `None` (hợp lệ), `str`/`list[str]` (lỗi), hoặc
    `dict` dạng `{"error": ..., "warning": ..., "data": {...}}`. Xem docstring `_RowContext`.

Nơi gọi (JS): alumglass/public/js/excel_io.js — 3 endpoint trong file này:
  alumglass.alumglass.excel_io.download_template / .upload / .export_xlsx
"""

from __future__ import unicode_literals

import datetime
import inspect
import re
from io import BytesIO

import frappe
import openpyxl
from frappe import _
from frappe.utils import cint, flt
from openpyxl.utils import get_column_letter

# --- Trần an toàn -----------------------------------------------------------
# File vài nghìn dòng thì nạp bình thường; vượt trần phải báo lỗi người-đọc-được thay vì treo
# request rồi trả về trang lỗi trắng. Trần kích thước còn phải chặn ở tầng nginx (client_max_body_size)
# — ở đây chặn để người dùng nhận được câu giải thích thay vì lỗi 413 của máy chủ.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UPLOAD_ROWS = 5000
MAX_EXPORT_ROWS = 20000
EXIST_CHUNK = 500
GUIDE_SHEET_TITLE = "Hướng dẫn"
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# `validator` do client gửi lên là một dotted path tuỳ ý — nếu nhận bừa thì endpoint thành công cụ
# gọi hàm bất kỳ (client tự chọn hàm nào cũng được gọi). Vì vậy chỉ nhận path nằm trong app này.
ALLOWED_VALIDATOR_PREFIXES = ("alumglass.",)
_DOTTED_PATH_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*(\.[A-Za-z_][A-Za-z0-9_]*)+$")

_SHEET_TITLE_RE = re.compile(r"[\\/\?\*\[\]:]")
_FILENAME_RE = re.compile(r"[^A-Za-z0-9._-]")


# ---------------------------------------------------------------------------
# Tiện ích dùng chung
# ---------------------------------------------------------------------------
def _as_json(value, default=None):
	"""Tham số đi từ JS sang luôn là chuỗi JSON (FormData) — hoặc đã là list/dict khi gọi nội bộ."""
	if value is None or value == "":
		return default
	if isinstance(value, (list, dict)):
		return value
	try:
		return frappe.parse_json(value)
	except Exception:
		frappe.throw(_("Dữ liệu gửi lên không hợp lệ. Vui lòng tải lại trang rồi thử lại."))


def _cell_text(value):
	"""Đổi giá trị 1 ô Excel thành chuỗi để validator làm việc.

	Ô ngày tháng openpyxl trả về datetime, còn số thì trả về float — nếu để nguyên `str()` thì
	ngày ra '2026-01-01 00:00:00' và số ra '1.0', cả hai đều không khớp với mã người dùng gõ tay.
	"""
	if value is None:
		return ""
	if isinstance(value, datetime.datetime):
		if value.time() == datetime.time.min:
			return value.strftime("%Y-%m-%d")
		return value.strftime("%Y-%m-%d %H:%M:%S")
	if isinstance(value, datetime.date):
		return value.strftime("%Y-%m-%d")
	if isinstance(value, bool):
		return "1" if value else "0"
	if isinstance(value, float) and value.is_integer():
		# Excel không phân biệt số nguyên với số thực; 1.0 phải đọc ra "1"
		return str(int(value))
	return str(value).strip()


def _header_key(value):
	"""So tiêu đề cột không phân biệt hoa/thường và khoảng trắng thừa hai đầu."""
	return _cell_text(value).lower()


def _normalize_columns(columns):
	"""Chuẩn hoá khai báo cột từ client; thiếu fieldname/label thì báo lỗi ngay (lỗi lập trình,
	không phải lỗi người dùng — nhưng vẫn phải có câu đọc được thay vì KeyError)."""
	columns = _as_json(columns, [])
	if not isinstance(columns, list) or not columns:
		frappe.throw(_("Chưa khai báo cột nào cho file Excel."))

	out = []
	seen = set()
	for index, col in enumerate(columns):
		if not isinstance(col, dict):
			frappe.throw(_("Khai báo cột thứ {0} không hợp lệ.").format(index + 1))
		fieldname = str(col.get("fieldname") or "").strip()
		if not fieldname:
			frappe.throw(_("Cột thứ {0} chưa có tên trường (fieldname).").format(index + 1))
		if fieldname in seen:
			frappe.throw(_("Cột '{0}' bị khai báo trùng.").format(fieldname))
		seen.add(fieldname)
		out.append({
			"fieldname": fieldname,
			"label": str(col.get("label") or fieldname).strip(),
			"fieldtype": str(col.get("fieldtype") or "Data").strip(),
			"options": (col.get("options") or "").strip(),
			"required": cint(col.get("required")),
			"width": cint(col.get("width")) or 16,
			"sample": _cell_text(col.get("sample")),
			"guide": str(col.get("guide") or "").strip(),
			# Cột chỉ để hiển thị/xuất (vd số tiền do validator tính) có thể tắt `export: 0`
			"export": 0 if col.get("export") in (0, "0", False) else 1,
		})
	return out


def _send_xlsx(content, filename):
	"""Trả file .xlsx về trình duyệt (tham số frappe.response giống các endpoint tải file sẵn có)."""
	frappe.response["filecontent"] = content
	frappe.response["filename"] = filename
	frappe.response["filetype"] = XLSX_MIMETYPE
	frappe.response["type"] = "binary"


def _safe_filename(name, suffix):
	"""Tên file tải về chỉ giữ ký tự ASCII an toàn.

	frappe.utils.response.as_binary đọc header Content-Disposition qua
	`filename.encode("utf-8").decode("unicode-escape")` — tên có dấu tiếng Việt đi qua đó sẽ thành
	chuỗi rác. Vì vậy tên file luôn được lọc về [A-Za-z0-9._-] trước khi trả.
	"""
	name = _FILENAME_RE.sub("_", str(name or "").strip()) or "export"
	if not name.lower().endswith(suffix):
		name = "{0}{1}".format(name, suffix)
	return name


def _sheet_title(name):
	"""Tên sheet Excel: tối đa 31 ký tự, không chứa \\ / ? * [ ] :"""
	return _SHEET_TITLE_RE.sub("-", str(name or "").strip())[:31] or "Sheet1"


def _workbook_bytes(workbook):
	fio = BytesIO()
	workbook.save(fio)
	return fio.getvalue()


# ---------------------------------------------------------------------------
# Hỏi DB theo lô — chống N+1
# ---------------------------------------------------------------------------
def get_existing(doctype, field, values, fields=None, extra_filters=None, chunk_size=EXIST_CHUNK):
	"""Tra DB cho nhiều giá trị trong MỘT lượt (chia lô 500), trả `{giá_trị: {tên_trường: giá_trị}}`.

	Dùng thay cho `frappe.db.exists`/`frappe.get_doc` trong vòng lặp: file vài nghìn dòng mà hỏi
	từng dòng thì mỗi dòng một query, chậm và dễ làm nghẽn. Chia lô để câu `IN (...)` không phình
	quá lớn (SQLite/MySQL đều có giới hạn tham số).
	"""
	if not values:
		return {}

	wanted = [fields] if isinstance(fields, str) else list(fields or [field])
	if field not in wanted:
		wanted.append(field)

	unique, seen = [], set()
	for value in values:
		key = _cell_text(value)
		if key and key not in seen:
			seen.add(key)
			unique.append(key)

	found = {}
	for start in range(0, len(unique), chunk_size):
		chunk = unique[start:start + chunk_size]
		filters = {field: ["in", chunk]}
		if extra_filters:
			# extra_filters KHÔNG được ghi đè chính `field` đang tra — ghi đè thì lô trả về sai
			for key, value in extra_filters.items():
				filters.setdefault(key, value)
		for row in frappe.db.get_all(doctype, filters=filters, fields=wanted) or []:
			key = _cell_text(row.get(field))
			if key:
				found[key] = row
	return found


class _Batch:
	"""Dữ liệu tra chung cho CẢ file, nạp một lần rồi dùng lại cho mọi dòng.

	Validator chạy từng dòng, nên nếu mỗi dòng tự gọi DB thì lại thành N+1. Mọi hàm ở đây đều tra
	theo **tập giá trị của cả cột trong file** (không phải giá trị của riêng dòng đang xét) nên dù
	dòng nào gọi trước thì cũng chỉ tốn đúng một lượt query cho mỗi (cột, danh sách trường).
	"""

	def __init__(self, doctype, rows):
		self.doctype = doctype
		self.rows = rows
		self._values = {}
		self._cache = {}
		self._queried = {}

	def values(self, fieldname):
		"""Các giá trị khác nhau của một cột trên toàn file (giữ thứ tự xuất hiện)."""
		if fieldname not in self._values:
			out, seen = [], set()
			for row in self.rows:
				value = _cell_text(row.get(fieldname))
				if value and value not in seen:
					seen.add(value)
					out.append(value)
			self._values[fieldname] = out
		return self._values[fieldname]

	def lookup(self, doctype, field, values, fields=None, extra_filters=None):
		"""Tra DB có cache; chỉ hỏi những giá trị chưa từng hỏi."""
		wanted = [fields] if isinstance(fields, str) else list(fields or [field])
		if field not in wanted:
			wanted.append(field)
		cache_key = (doctype, field, tuple(wanted), frappe.as_json(extra_filters or {}))
		cached = self._cache.setdefault(cache_key, {})
		# Nhớ cả giá trị ĐÃ HỎI nhưng không có trong DB — nếu chỉ nhớ giá trị tìm thấy thì mã không
		# tồn tại sẽ bị hỏi lại ở từng dòng, đúng cái N+1 cần tránh.
		queried = self._queried.setdefault(cache_key, set())

		missing = []
		for value in values:
			key = _cell_text(value)
			if key and key not in queried:
				queried.add(key)
				missing.append(key)

		if missing:
			cached.update(get_existing(
				doctype, field, missing, fields=wanted, extra_filters=extra_filters
			))
		return cached

	def existing(self, column, field=None, fields=None, extra_filters=None):
		"""Tra cả cột `column` của file trong 1 lượt rồi trả `{giá_trị: {trường: giá_trị}}`.

		`column` là tên cột trong file, `field` là tên trường trong DB để so — hai tên này thường
		KHÁC nhau (cột Excel 'purchase_receipt' ứng với trường 'name' của Purchase Receipt), nên
		phải khai tường minh; mặc định `field = column` cho trường hợp trùng tên.
		"""
		field = field or column
		return self.lookup(self.doctype, field, self.values(column), fields, extra_filters)


class _RowContext:
	"""Ngữ cảnh truyền vào validator cho mỗi dòng.

	Thuộc tính:
	  - `ctx.doctype`    : doctype đã khai trên dialog (dùng để kiểm quyền ở `upload`)
	  - `ctx.line`       : số dòng trong file Excel (dòng 2 là dòng dữ liệu đầu tiên)
	  - `ctx.key`        : tên trường của cột khoá (rỗng nếu dialog không khai)
	  - `ctx.extra_args` : dict client gửi kèm (vd `{company, supplier}` đang chọn trên dialog)
	  - `ctx.batch`      : `_Batch` — tra DB theo lô, xem docstring `_Batch`

	Ví dụ validator:

	    def validate_receipt_rows(row, ctx):
	        name = row["purchase_receipt"]
	        # hỏi 1 lượt cho cả cột của file rồi tra trong bộ nhớ, KHÔNG hỏi DB trong vòng lặp
	        existing = ctx.batch.existing("purchase_receipt", field="name",
	                                      fields=["name", "docstatus"])
	        if name not in existing:
	            return "'{0}' — không tìm thấy Purchase Receipt này".format(name)
	        if existing[name].get("docstatus") == 2:
	            return "'{0}' — Purchase Receipt này đã bị huỷ, không dùng được".format(name)

	Validator trả về:
	  - `None` / `""` / `[]`  -> dòng hợp lệ
	  - `"câu lỗi"`           -> dòng lỗi (KHÔNG tự thêm "Dòng N:" — thư viện đã thêm)
	  - `["lỗi 1", "lỗi 2"]`  -> nhiều lỗi cho cùng dòng
	  - `{"error": ..., "warning": ..., "data": {...}}` -> lỗi/cảnh báo/giá trị bổ sung cho dòng
	    (khoá `data` dùng để validator đắp thêm cột mà file không có, vd số tiền, ngày, nhà cung cấp)
	"""

	def __init__(self, doctype, line, key, extra_args, batch):
		self.doctype = doctype
		self.line = line
		self.key = key
		self.extra_args = extra_args or {}
		self.batch = batch


# ---------------------------------------------------------------------------
# Đọc file
# ---------------------------------------------------------------------------
def _get_uploaded_file():
	"""Lấy file từ request, chặn trước khi đọc: sai đuôi / quá lớn."""
	if "file" not in (frappe.request.files or {}):
		frappe.throw(_("Vui lòng chọn file Excel để tải lên"))

	file = frappe.request.files["file"]
	# So đuôi file KHÔNG phân biệt hoa/thường — người dùng hay có file .XLSX, chặn nhầm là lỗi vô cớ
	if not (file.filename or "").lower().endswith(".xlsx"):
		frappe.throw(_("File không đúng định dạng. Vui lòng chọn file Excel (.xlsx)"))

	size = 0
	try:
		file.stream.seek(0, 2)
		size = file.stream.tell()
		file.stream.seek(0)
	except Exception:
		# Stream không seek được thì đành bỏ qua bước đo — openpyxl vẫn đọc được
		size = 0

	if size and size > MAX_UPLOAD_BYTES:
		frappe.throw(_(
			"File có dung lượng {0} MB, vượt quá giới hạn {1} MB mỗi lần tải lên. "
			"Vui lòng chia nhỏ file rồi tải lên lại."
		).format(round(size / 1024.0 / 1024.0, 1), int(MAX_UPLOAD_BYTES / 1024 / 1024)))

	return file


def _read_rows(file, columns):
	"""Đọc file Excel thành danh sách dòng `{fieldname: giá_trị}` (kèm `_line`).

	Header phải nằm ở dòng 1. Ô A1 phải đúng nhãn của cột đầu tiên trong `columns` — kiểm chặt
	theo bản Purchase Invoice: chỉ kiểm "ô không rỗng đầu tiên" thì file đặt tiêu đề ở cột B vẫn
	qua được, rồi đọc ra toàn rỗng và người dùng không hiểu vì sao không nạp được gì.
	"""
	try:
		workbook = openpyxl.load_workbook(file.stream, read_only=True, data_only=True)
	except Exception:
		frappe.throw(_("Chưa đọc được file Excel. Vui lòng kiểm tra file đúng định dạng .xlsx rồi thử lại."))

	try:
		sheet = workbook.active
		if not sheet:
			frappe.throw(_("File Excel không có dữ liệu"))

		header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
		if not header_row:
			frappe.throw(_("File Excel không có dữ liệu"))

		first_label = columns[0]["label"]
		first_header = _header_key(header_row[0] if header_row else None)
		if first_header != _header_key(first_label):
			found = _cell_text(header_row[0]) if header_row else ""
			if found:
				frappe.throw(_(
					"File Excel không đúng định dạng. Ô A1 (dòng đầu, cột A) phải là tiêu đề '{0}' "
					"nhưng đang là '{1}'."
				).format(first_label, found))
			frappe.throw(_(
				"File Excel không đúng định dạng. Ô A1 (dòng đầu, cột A) phải là tiêu đề '{0}'."
			).format(first_label))

		# Map cột theo TÊN ở header (không theo vị trí) — người dùng xoá/đổi thứ tự cột phụ
		# vẫn nạp đúng. Cột bắt buộc mà không thấy tiêu đề thì báo luôn, không để nạp ra rỗng.
		index_by_key = {}
		for index, value in enumerate(header_row):
			key = _header_key(value)
			if key and key not in index_by_key:
				index_by_key[key] = index

		column_index = {}
		missing_headers = []
		for col in columns:
			index = index_by_key.get(_header_key(col["label"]))
			if index is None:
				index = index_by_key.get(_header_key(col["fieldname"]))
			if index is None:
				if col["required"]:
					missing_headers.append(col["label"])
				continue
			column_index[col["fieldname"]] = index

		if missing_headers:
			frappe.throw(_(
				"File Excel thiếu cột bắt buộc: {0}. Vui lòng tải file mẫu để có đúng tiêu đề cột."
			).format(", ".join("'{0}'".format(label) for label in missing_headers)))

		rows = []
		for line, raw in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
			row = {"_line": line}
			has_value = False
			for col in columns:
				index = column_index.get(col["fieldname"])
				value = _cell_text(raw[index]) if index is not None and index < len(raw) else ""
				row[col["fieldname"]] = value
				if value:
					has_value = True
			if has_value:
				rows.append(row)
			if len(rows) > MAX_UPLOAD_ROWS:
				frappe.throw(_(
					"File có nhiều hơn {0} dòng dữ liệu, vượt quá giới hạn mỗi lần tải lên. "
					"Vui lòng chia nhỏ file rồi tải lên lại."
				).format(MAX_UPLOAD_ROWS))
	finally:
		try:
			workbook.close()
		except Exception:
			pass

	return rows


# ---------------------------------------------------------------------------
# Validator
# ---------------------------------------------------------------------------
def _resolve_validator(path):
	path = str(path or "").strip()
	if not path:
		return None
	if not _DOTTED_PATH_RE.match(path):
		frappe.throw(_("Hàm kiểm tra dữ liệu không hợp lệ."))
	if not path.startswith(ALLOWED_VALIDATOR_PREFIXES):
		frappe.throw(_("Hàm kiểm tra dữ liệu không hợp lệ."))
	try:
		func = frappe.get_attr(path)
	except Exception:
		frappe.throw(_("Không tìm thấy hàm kiểm tra dữ liệu '{0}'.").format(path))
	if inspect.isclass(func) or not callable(func):
		frappe.throw(_("Hàm kiểm tra dữ liệu không hợp lệ."))
	return func


def _call_validator(func, row, ctx):
	"""Gọi validator 1 hoặc 2 tham số (`validate(row)` / `validate(row, ctx)`) — dialog cũ chỉ cần
	đọc dữ liệu trong dòng thì không phải nhận thêm tham số không dùng."""
	try:
		params = inspect.signature(func).parameters.values()
		positional = [
			p for p in params
			if p.kind in (p.POSITIONAL_ONLY, p.POSITIONAL_OR_KEYWORD)
		]
		varargs = any(p.kind == p.VAR_POSITIONAL for p in params)
	except (TypeError, ValueError):
		positional, varargs = [None, None], False

	if varargs or len(positional) >= 2:
		return func(row, ctx)
	return func(row)


def _normalize_validator_result(result):
	"""Đưa kết quả validator về `(errors, warnings, data)`."""
	if result is None or result == "" or result == []:
		return [], [], {}
	if isinstance(result, str):
		return [result], [], {}
	if isinstance(result, (list, tuple)):
		return [str(item) for item in result if item], [], {}
	if isinstance(result, dict):
		errors = result.get("error") or result.get("errors") or []
		warnings = result.get("warning") or result.get("warnings") or []
		if isinstance(errors, str):
			errors = [errors]
		if isinstance(warnings, str):
			warnings = [warnings]
		data = result.get("data") or result.get("row") or {}
		if not isinstance(data, dict):
			data = {}
		return [str(item) for item in errors if item], [str(item) for item in warnings if item], data
	return [str(result)], [], {}


# ---------------------------------------------------------------------------
# Endpoint: file mẫu
# ---------------------------------------------------------------------------
@frappe.whitelist()
def download_template(doctype=None, columns=None, filename=None, title=None, guide=None):
	"""Sinh file Excel mẫu: dòng 1 là tiêu đề cột, dòng 2 là dòng mẫu, kèm sheet "Hướng dẫn"."""
	columns = _normalize_columns(columns)

	if doctype and not frappe.has_permission(doctype, "read"):
		frappe.throw(_("Bạn không có quyền xem {0} nên không tải được file mẫu").format(_(doctype)))

	workbook = openpyxl.Workbook()
	sheet = workbook.active
	sheet.title = _sheet_title(title or doctype or "Sheet1")
	sheet.append([col["label"] for col in columns])

	# Dòng mẫu: chỉ ghi khi có cột khai `sample` — cột để trống thì dòng mẫu không mang thông tin gì,
	# còn ghi giá trị bịa vào thì người dùng quên xoá sẽ nạp nhầm và bị báo lỗi khó hiểu.
	if any(col["sample"] for col in columns):
		# Ô trống ghi bằng None (không ghi cell) thay vì chuỗi rỗng — openpyxl ghi chuỗi rỗng thành
		# `<c t="inlineStr">` thiếu phần `<is>`, một số bản Excel kêu "file cần sửa" khi mở
		sheet.append([col["sample"] or None for col in columns])

	for index, col in enumerate(columns, start=1):
		sheet.column_dimensions[get_column_letter(index)].width = col["width"]

	_append_guide_sheet(workbook, columns, guide)
	_send_xlsx(_workbook_bytes(workbook), _safe_filename(filename or doctype or "template", "_upload_template.xlsx"))


def _append_guide_sheet(workbook, columns, guide):
	"""Sheet "Hướng dẫn" mô tả từng cột — để người dùng mở file là biết phải điền gì."""
	sheet = workbook.create_sheet(GUIDE_SHEET_TITLE)
	sheet.column_dimensions["A"].width = 28
	sheet.column_dimensions["B"].width = 14
	sheet.column_dimensions["C"].width = 16
	sheet.column_dimensions["D"].width = 70

	sheet.append([str(guide) if guide else "Cách dùng file này:"])
	sheet.append([])
	sheet.append(["Cột", "Bắt buộc", "Kiểu dữ liệu", "Mô tả"])

	for col in columns:
		sheet.append([
			col["label"],
			"Bắt buộc" if col["required"] else "Không bắt buộc",
			col["fieldtype"] + (" ({0})".format(col["options"]) if col["options"] else ""),
			col["guide"] or _default_column_guide(col),
		])

	sheet.append([])
	sheet.append(["Lưu ý", "", "", "Chỉ nhận file .xlsx. Dòng 1 là tiêu đề cột — không sửa tiêu đề."])
	sheet.append(["", "", "", "Nạp file sẽ thay thế toàn bộ nội dung bảng trên dialog."])
	sheet.append(["", "", "", "Dòng sai được báo lại kèm số dòng trong file để sửa rồi tải lên lại."])


def _default_column_guide(col):
	if col["fieldtype"] == "Link" and col["options"]:
		return _("Mã {0} đã có trong hệ thống").format(col["options"])
	return _("Giá trị của cột {0}").format(col["label"])


# ---------------------------------------------------------------------------
# Endpoint: nạp file
# ---------------------------------------------------------------------------
@frappe.whitelist()
def upload(doctype=None, columns=None, validator=None, key=None, extra_args=None):
	"""Đọc file Excel -> map cột theo tên ở header -> kiểm tra -> trả `{results, errors, warnings}`.

	Phần kiểm tra chung (tiêu đề cột, cột bắt buộc, dòng trùng khoá) do thư viện làm; luật nghiệp
	vụ do `validator` của từng dialog làm. Không khai `validator` thì chạy chế độ raw: mọi dòng đọc
	được đều nằm trong `results`.
	"""
	columns = _normalize_columns(columns)
	extra_args = _as_json(extra_args, {}) or {}
	key = str(key or "").strip()

	if not doctype:
		frappe.throw(_("Chưa khai báo loại chứng từ cho lần nạp này."))

	# Kiểm quyền TRƯỚC khi đọc file — không đọc nội dung file của người không có quyền
	if not frappe.has_permission(doctype, "read"):
		frappe.throw(_("Bạn không có quyền xem {0} nên không nạp được danh sách này").format(_(doctype)))

	if key and key not in [col["fieldname"] for col in columns]:
		frappe.throw(_("Cột khoá '{0}' chưa được khai báo trong danh sách cột.").format(key))

	validate = _resolve_validator(validator)
	file = _get_uploaded_file()
	rows = _read_rows(file, columns)

	errors = []      # [(số dòng, câu lỗi)] — gom rồi sắp theo số dòng ở cuối
	warnings = []
	results = []

	# --- Dòng trùng khoá: giữ dòng đầu, báo rõ dòng trùng với dòng nào -------------------
	duplicates = []
	first_seen = {}
	if key:
		kept = []
		for row in rows:
			value = _cell_text(row.get(key))
			if value in first_seen:
				duplicates.append((row["_line"], value, first_seen[value]))
				continue
			first_seen[value] = row["_line"]
			kept.append(row)
		rows = kept

	if not rows:
		return {
			"results": [],
			"errors": [
				_("Dòng {0}: '{1}' — trùng với dòng {2}, đã bỏ qua").format(line, value, first)
				for line, value, first in duplicates
			],
			"warnings": [],
			"total": 0,
		}

	batch = _Batch(doctype, rows)

	for row in rows:
		line = row["_line"]
		field_row = {k: v for k, v in row.items() if k != "_line"}

		# Cột bắt buộc để trống thì báo tại chỗ, không gọi validator (validator thường giả định
		# đã có giá trị, để nó tự kiểm thì mỗi dialog phải lặp lại cùng một câu kiểm tra)
		missing = [
			col["label"] for col in columns
			if col["required"] and not _cell_text(field_row.get(col["fieldname"]))
		]
		if missing:
			errors.append((line, _("chưa nhập {0}").format(", ".join(missing))))
			continue

		row_errors, row_warnings, data = [], [], {}
		if validate:
			ctx = _RowContext(doctype, line, key, extra_args, batch)
			row_errors, row_warnings, data = _normalize_validator_result(
				_call_validator(validate, field_row, ctx)
			)

		errors.extend((line, message) for message in row_errors)
		warnings.extend((line, message) for message in row_warnings)

		if not row_errors:
			final = dict(field_row)
			final.update(data)
			results.append(final)

	for line, value, first in duplicates:
		errors.append((line, _("'{0}' — trùng với dòng {1}, đã bỏ qua").format(value, first)))

	# Sắp theo số dòng: người dùng đọc file từ trên xuống là gặp lỗi theo đúng thứ tự đó
	errors.sort(key=lambda pair: pair[0])
	warnings.sort(key=lambda pair: pair[0])

	return {
		"results": results,
		"errors": [_("Dòng {0}: {1}").format(line, message) for line, message in errors],
		"warnings": [_("Dòng {0}: {1}").format(line, message) for line, message in warnings],
		"total": len(rows) + len(duplicates),
	}


# ---------------------------------------------------------------------------
# Endpoint: xuất Excel
# ---------------------------------------------------------------------------
@frappe.whitelist()
def export_xlsx(doctype=None, columns=None, rows=None, filename=None):
	"""Xuất dữ liệu đang có trên bảng của dialog ra file .xlsx."""
	columns = _normalize_columns(columns)
	rows = _as_json(rows, [])

	if doctype and not frappe.has_permission(doctype, "read"):
		frappe.throw(_("Bạn không có quyền xem {0} nên không xuất được dữ liệu").format(_(doctype)))

	if not isinstance(rows, list) or not rows:
		frappe.throw(_("Chưa có dòng nào để xuất."))

	if len(rows) > MAX_EXPORT_ROWS:
		frappe.throw(_(
			"Bảng có {0} dòng, vượt quá giới hạn {1} dòng mỗi lần xuất. Vui lòng lọc bớt rồi xuất lại."
		).format(len(rows), MAX_EXPORT_ROWS))

	export_columns = [col for col in columns if col["export"]]
	if not export_columns:
		frappe.throw(_("Chưa khai báo cột nào để xuất."))

	workbook = openpyxl.Workbook()
	sheet = workbook.active
	sheet.title = _sheet_title(doctype or "Data")
	sheet.append([col["label"] for col in export_columns])

	# Giá trị đi từ bảng trên màn hình ra file, tức là do người dùng nhập — có thể bắt đầu bằng "=".
	# openpyxl ghi chuỗi bắt đầu bằng "=" thành ô CÔNG THỨC thật trong file (đã kiểm chứng:
	# `<c r="A1"><f>1+1</f></c>`), mở bằng Excel là công thức chạy. Ép data_type="s" để ô được ghi
	# dạng chuỗi literal, giữ nguyên nội dung mà không phải thêm dấu nháy vào dữ liệu.
	for row in rows:
		if not isinstance(row, dict):
			continue
		values = []
		for col in export_columns:
			value = row.get(col["fieldname"])
			if value is None or value == "":
				# Ô trống ghi bằng None; chuỗi rỗng bị openpyxl ghi thành `<c t="inlineStr">` thiếu `<is>`
				values.append(None)
				continue
			if isinstance(value, (datetime.datetime, datetime.date)):
				values.append(_cell_text(value))
				continue
			if isinstance(value, bool):
				values.append("1" if value else "0")
				continue
			if isinstance(value, (int, float)):
				values.append(value)
				continue
			values.append(str(value))

		sheet.append(values)
		for index, value in enumerate(values, start=1):
			if isinstance(value, str):
				sheet.cell(row=sheet.max_row, column=index).data_type = "s"

	for index, col in enumerate(export_columns, start=1):
		sheet.column_dimensions[get_column_letter(index)].width = col["width"]

	_send_xlsx(_workbook_bytes(workbook), _safe_filename(filename or doctype or "export", ".xlsx"))
